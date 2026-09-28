"""Generate explainable research charts for thesis / viva.

Outputs PNG charts under data/exports/analysis/charts/ and a short Word
pack with the charts embedded for supervisor demos.
"""

from __future__ import annotations

import csv
import sqlite3
from collections import Counter
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
ANALYSIS = ROOT / "data" / "exports" / "analysis"
CHARTS = ANALYSIS / "charts"
DOCX = ROOT / "Research_Charts_Explainable_IM2021049.docx"
DB = ROOT / "data" / "smartphones.db"

# Teal / slate palette (matches UI; avoids purple bias)
TEAL = "#0f766e"
TEAL_SOFT = "#5eead4"
SLATE = "#334155"
POS = "#047857"
NEU = "#b45309"
NEG = "#b91c1c"
ACCENTS = ["#0f766e", "#0d9488", "#14b8a6", "#2dd4bf", "#5eead4", "#99f6e4", "#ccfbf1", "#f0fdfa"]


def font(run, *, bold=False, size=11):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def style_ax(ax, title: str):
    ax.set_title(title, fontsize=12, fontweight="bold", color=SLATE, pad=10)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(colors=SLATE)
    ax.yaxis.label.set_color(SLATE)
    ax.xaxis.label.set_color(SLATE)


def save(fig, name: str) -> Path:
    CHARTS.mkdir(parents=True, exist_ok=True)
    path = CHARTS / name
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def chart_brand_distribution() -> Path:
    # Prefer live DB counts so carrier/manual cleanups show up immediately.
    brands_counts: list[tuple[str, int]] = []
    if DB.is_file():
        conn = sqlite3.connect(DB)
        brands_counts = conn.execute(
            "SELECT brand, COUNT(*) AS n FROM smartphones "
            "WHERE brand IS NOT NULL AND TRIM(brand) != '' "
            "GROUP BY brand ORDER BY n DESC, brand"
        ).fetchall()
        conn.close()
    if not brands_counts:
        rows = read_csv(ANALYSIS / "01_brand_coverage.csv")
        brands_counts = [(r["Brand"], int(r["Phones"])) for r in rows]

    brands = [b for b, _ in brands_counts]
    counts = [c for _, c in brands_counts]
    # collapse tiny brands
    major, other = [], 0
    for b, c in zip(brands, counts):
        if c >= 6:
            major.append((b, c))
        else:
            other += c
    if other:
        major.append(("Other", other))
    labels = [m[0] for m in major]
    values = [m[1] for m in major]

    fig, axes = plt.subplots(1, 2, figsize=(10.5, 4.2))
    axes[0].barh(labels[::-1], values[::-1], color=TEAL)
    style_ax(axes[0], "Smartphone count by brand")
    axes[0].set_xlabel("Number of phones in corpus")
    for i, v in enumerate(values[::-1]):
        axes[0].text(v + 1, i, str(v), va="center", fontsize=9, color=SLATE)

    colors = ACCENTS[: len(labels)]
    wedges, texts, autotexts = axes[1].pie(
        values,
        labels=labels,
        autopct=lambda p: f"{p:.0f}%" if p >= 3 else "",
        colors=colors,
        startangle=90,
        textprops={"color": SLATE, "fontsize": 9},
    )
    for t in autotexts:
        t.set_color("white")
        t.set_fontsize(8)
        t.set_fontweight("bold")
    style_ax(axes[1], "Brand share of corpus")
    fig.suptitle(
        "Figure A: Distribution of mobile phones by brand (as-built corpus)",
        fontsize=11,
        color=SLATE,
        y=1.02,
    )
    return save(fig, "01_phone_brand_distribution.png")


def chart_reviews_per_phone() -> Path:
    conn = sqlite3.connect(DB)
    counts = [
        r[0]
        for r in conn.execute(
            """
            select count(*) from reviews
            where excluded_reason is null or excluded_reason=''
            group by smartphone_id
            """
        )
    ]
    conn.close()
    fig, ax = plt.subplots(figsize=(8, 4.2))
    bins = range(0, max(counts) + 2)
    ax.hist(counts, bins=bins, color=TEAL, edgecolor="white", align="left")
    style_ax(ax, "How many kept reviews does each phone have?")
    ax.set_xlabel("Kept reviews per phone (after preprocessing)")
    ax.set_ylabel("Number of phones")
    mean = sum(counts) / len(counts) if counts else 0
    ax.axvline(mean, color=NEG, linestyle="--", linewidth=1.5, label=f"Mean ≈ {mean:.1f}")
    ax.legend(frameon=False)
    fig.suptitle(
        "Figure B: Review-volume distribution across phones",
        fontsize=11,
        color=SLATE,
        y=1.02,
    )
    return save(fig, "02_reviews_per_phone_hist.png")


def chart_aspect_sentiment() -> Path:
    rows = [r for r in read_csv(ANALYSIS / "02_per_aspect_sentiment.csv") if r["Aspect"] != "Total"]
    aspects = [r["Aspect"] for r in rows]
    pos = [int(r["Positive"]) for r in rows]
    neu = [int(r["Neutral"]) for r in rows]
    neg = [int(r["Negative"]) for r in rows]
    totals = [int(r["Total"]) for r in rows]

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5))

    x = range(len(aspects))
    axes[0].bar(x, pos, label="Positive", color=POS)
    axes[0].bar(x, neu, bottom=pos, label="Neutral", color=NEU)
    axes[0].bar(
        x,
        neg,
        bottom=[p + n for p, n in zip(pos, neu)],
        label="Negative",
        color=NEG,
    )
    axes[0].set_xticks(list(x))
    axes[0].set_xticklabels(aspects, rotation=20, ha="right")
    style_ax(axes[0], "Sentiment mix per smartphone aspect")
    axes[0].set_ylabel("Mention count")
    axes[0].legend(frameon=False, fontsize=8)

    axes[1].barh(aspects[::-1], totals[::-1], color=TEAL)
    style_ax(axes[1], "Which features are discussed most?")
    axes[1].set_xlabel("Total ABSA mentions")
    for i, v in enumerate(totals[::-1]):
        pct = 100 * v / sum(totals)
        axes[1].text(v + 30, i, f"{v:,} ({pct:.0f}%)", va="center", fontsize=8, color=SLATE)

    fig.suptitle(
        "Figure C: Aspect-Based Sentiment Analysis — distribution of review opinions",
        fontsize=11,
        color=SLATE,
        y=1.02,
    )
    return save(fig, "03_aspect_sentiment_distribution.png")


def chart_overall_sentiment() -> Path:
    rows = read_csv(ANALYSIS / "02_per_aspect_sentiment.csv")
    total = next(r for r in rows if r["Aspect"] == "Total")
    labels = ["Positive", "Neutral", "Negative"]
    values = [int(total["Positive"]), int(total["Neutral"]), int(total["Negative"])]
    colors = [POS, NEU, NEG]
    fig, ax = plt.subplots(figsize=(5.5, 4.5))
    wedges, texts, autotexts = ax.pie(
        values,
        labels=labels,
        autopct="%1.1f%%",
        colors=colors,
        startangle=90,
        explode=(0.02, 0.02, 0.05),
        textprops={"color": SLATE},
    )
    for t in autotexts:
        t.set_color("white")
        t.set_fontweight("bold")
    style_ax(ax, "Overall sentiment of all aspect mentions\n(n = 13,914)")
    fig.suptitle("Figure D: Overall ABSA sentiment mix", fontsize=11, color=SLATE, y=1.02)
    return save(fig, "04_overall_sentiment_pie.png")


def chart_sample_phone_scores() -> Path:
    conn = sqlite3.connect(DB)
    phones = conn.execute(
        """
        select s.id, coalesce(s.canonical_name, s.raw_title), s.source_product_id,
               sum(a.mention_count) m
        from aspect_scores a join smartphones s on s.id=a.smartphone_id
        group by s.id having count(*)=6
        order by m desc limit 2
        """
    ).fetchall()
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharey=True)
    order = ["battery", "camera", "display", "performance", "design", "price"]
    for ax, (pid, name, asin, _m) in zip(axes, phones):
        rows = {
            r[0]: r
            for r in conn.execute(
                """
                select aspect, score, mention_count, confidence
                from aspect_scores where smartphone_id=?
                """,
                (pid,),
            )
        }
        labels = [a.title() for a in order]
        scores = [float(rows[a][1]) for a in order]
        conf = [float(rows[a][3]) for a in order]
        bars = ax.barh(labels[::-1], scores[::-1], color=TEAL)
        for bar, c in zip(bars, conf[::-1]):
            # lighter bars for low confidence
            if c < 0.6:
                bar.set_alpha(0.45)
        style_ax(ax, f"{name[:28]}\n({asin})")
        ax.set_xlim(0, 1)
        ax.set_xlabel("Feature score (0–1, shrunk)")
        for i, (s, a) in enumerate(zip(scores[::-1], order[::-1])):
            m = rows[a][2]
            ax.text(min(s + 0.02, 0.92), i, f"{s:.2f} · {m} mentions", va="center", fontsize=7, color=SLATE)
    conn.close()
    fig.suptitle(
        "Figure E: Explainable feature-score profiles (faded bars = lower confidence)",
        fontsize=11,
        color=SLATE,
        y=1.03,
    )
    return save(fig, "05_sample_phone_feature_scores.png")


def chart_agreement_by_stars() -> Path:
    path = ANALYSIS / "05_star_text_agreement_by_rating.csv"
    if not path.exists():
        return Path()
    rows = read_csv(path)
    stars = [r["stars"] for r in rows]
    rates = [float(r["agreement_rate"]) * 100 for r in rows]
    reviews = [int(r["reviews"]) for r in rows]
    fig, ax = plt.subplots(figsize=(7.5, 4))
    bars = ax.bar(stars, rates, color=TEAL, width=0.65)
    style_ax(ax, "Does ABSA polarity agree with Amazon star ratings?")
    ax.set_xlabel("Review star rating")
    ax.set_ylabel("Agreement rate (%)")
    ax.set_ylim(0, 100)
    for bar, n, r in zip(bars, reviews, rates):
        ax.text(bar.get_x() + bar.get_width() / 2, r + 2, f"n={n}", ha="center", fontsize=8, color=SLATE)
    fig.suptitle(
        "Figure F: Star–text agreement by rating (ABSA sanity check)",
        fontsize=11,
        color=SLATE,
        y=1.02,
    )
    return save(fig, "06_star_text_agreement_by_rating.png")


def chart_ranking_profiles() -> Path:
    profiles = [
        ("04_top5_balanced.csv", "Balanced"),
        ("04_top5_camera-focused.csv", "Camera"),
        ("04_top5_battery-focused.csv", "Battery"),
        ("04_top5_price-focused.csv", "Price"),
    ]
    # collect Top-1 score per profile + phone name
    names, scores, labels = [], [], []
    for fname, label in profiles:
        path = ANALYSIS / fname
        if not path.exists():
            continue
        rows = read_csv(path)
        if not rows:
            continue
        top = rows[0]
        labels.append(label)
        names.append(top["phone"][:22])
        scores.append(float(top["score"]))

    fig, ax = plt.subplots(figsize=(8.5, 4.2))
    bars = ax.bar(labels, scores, color=[TEAL, "#0d9488", "#14b8a6", "#5eead4"][: len(labels)])
    style_ax(ax, "Top-1 recommendation changes with user priorities")
    ax.set_ylabel("Top-1 weighted score")
    ax.set_ylim(0, 1)
    for bar, name, s in zip(bars, names, scores):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            s + 0.03,
            name,
            ha="center",
            va="bottom",
            fontsize=8,
            color=SLATE,
            rotation=15,
        )
    fig.suptitle(
        "Figure G: Explainable ranking — different priorities pick different Top-1 phones",
        fontsize=11,
        color=SLATE,
        y=1.02,
    )
    return save(fig, "07_ranking_priority_top1.png")


def chart_proposed_vs_baseline() -> Path:
    path = ANALYSIS / "04_proposed_vs_star_baseline_top10.csv"
    if not path.exists():
        return Path()
    rows = read_csv(path)[:5]
    ranks = [r["rank"] for r in rows]
    prop = [r["proposed_phone"][:18] for r in rows]
    base = [r["baseline_phone"][:18] for r in rows]

    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.axis("off")
    table_data = [["Rank", "Proposed (feature weights)", "Amazon stars baseline"]]
    for r, p, b in zip(ranks, prop, base):
        table_data.append([r, p, b])
    table = ax.table(
        cellText=table_data,
        loc="center",
        cellLoc="left",
        colWidths=[0.1, 0.45, 0.45],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(9)
    table.scale(1, 1.6)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor("#cbd5e1")
        if row == 0:
            cell.set_facecolor(TEAL)
            cell.get_text().set_color("white")
            cell.get_text().set_fontweight("bold")
        else:
            cell.set_facecolor("#f8fafc" if row % 2 == 0 else "white")
            cell.get_text().set_color(SLATE)
    ax.set_title(
        "Figure H: Proposed ranking vs Amazon star baseline (Top-5, balanced weights)\n"
        "Lists differ — feature priorities are not the same as popularity stars",
        fontsize=11,
        fontweight="bold",
        color=SLATE,
        pad=12,
    )
    return save(fig, "08_proposed_vs_baseline_table.png")


def build_docx(paths: list[Path]):
    doc = Document()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Explainable Research Charts — IM/2021/049")
    font(r, bold=True, size=16)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(
        s.add_run(
            "Visual figures for Chapter 5 / viva. Use these to explain phone distribution, "
            "ABSA results, feature scores, and ranking behaviour to your supervisor."
        ),
        size=11,
    )

    captions = {
        "01_phone_brand_distribution.png": (
            "Figure A — Phone distribution by brand. Shows the corpus is Apple-heavy; "
            "mention this as a limitation and coverage fact."
        ),
        "02_reviews_per_phone_hist.png": (
            "Figure B — Reviews per phone. Mean ≈ 15 kept reviews; not 200/phone — report honestly."
        ),
        "03_aspect_sentiment_distribution.png": (
            "Figure C — Aspect & sentiment distribution. Battery is discussed most; all six features appear."
        ),
        "04_overall_sentiment_pie.png": (
            "Figure D — Overall sentiment mix across 13,914 aspect mentions."
        ),
        "05_sample_phone_feature_scores.png": (
            "Figure E — Explainable feature scores for two phones. Faded bars = lower confidence "
            "(fewer mentions)."
        ),
        "06_star_text_agreement_by_rating.png": (
            "Figure F — ABSA vs Amazon stars by rating bucket (sanity check, not gold accuracy)."
        ),
        "07_ranking_priority_top1.png": (
            "Figure G — Changing user priorities changes the Top-1 recommendation."
        ),
        "08_proposed_vs_baseline_table.png": (
            "Figure H — Proposed method vs Amazon star baseline Top-5 (different lists)."
        ),
    }

    for path in paths:
        if not path or not path.exists():
            continue
        doc.add_paragraph()
        p = doc.add_paragraph()
        font(p.add_run(captions.get(path.name, path.stem)), bold=True, size=11)
        doc.add_picture(str(path), width=Inches(6.2))
        note = doc.add_paragraph()
        font(note.add_run(f"File: data/exports/analysis/charts/{path.name}"), size=9)

    viva = doc.add_paragraph()
    font(viva.add_run("Viva tip"), bold=True, size=12)
    tips = doc.add_paragraph()
    font(
        tips.add_run(
            "Open this document (or the PNG folder) when asked: "
            "“Show me the data distribution”, “How do you know ABSA works?”, "
            "“Do user priorities matter?”, or “How is this better than star ratings?”"
        ),
        size=11,
    )
    doc.save(DOCX)
    print(f"Wrote {DOCX}")

    # Sync into the web UI static folder so Evaluation tab can serve them.
    try:
        from app.services.eval_charts import sync_charts_to_static

        n = sync_charts_to_static()
        print(f"Synced {n} chart(s) to app/static/charts for /ui/ Evaluation tab")
    except Exception as exc:
        print("Static sync skipped:", exc)


def main():
    paths = [
        chart_brand_distribution(),
        chart_reviews_per_phone(),
        chart_aspect_sentiment(),
        chart_overall_sentiment(),
        chart_sample_phone_scores(),
        chart_agreement_by_stars(),
        chart_ranking_profiles(),
        chart_proposed_vs_baseline(),
    ]
    build_docx([p for p in paths if p])
    print("Charts saved to", CHARTS)


if __name__ == "__main__":
    main()
