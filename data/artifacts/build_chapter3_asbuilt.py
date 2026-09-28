"""Rewrite Chapter 3 Methodology to match the as-built system."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[2]
    / "Chapter3_Methodology_AsBuilt_IM2021049.docx"
)


def font(run, *, bold=False, italic=False, size=12):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")


def h(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        font(run, bold=True, size=14 if level == 1 else (13 if level == 2 else 12))


def p(doc, text, *, italic=False, indent=True, after=8):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Inches(0.5) if indent else Inches(0)
    run = para.add_run(text)
    font(run, italic=italic, size=12)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.clear()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    para.paragraph_format.first_line_indent = Inches(0)
    run = para.add_run(text)
    font(run, size=12)


def table(doc, headers, rows, caption=None):
    if caption:
        p(doc, caption, italic=True, indent=False, after=6)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        font(cell.paragraphs[0].add_run(head), bold=True, size=10)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = ""
            font(cell.paragraphs[0].add_run(str(val)), size=10)
    doc.add_paragraph()


def note(doc, text):
    p(doc, text, italic=True, indent=False, after=10)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("CHAPTER 3: METHODOLOGY")
    font(r, bold=True, size=16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "As-built methodology for the smartphone ABSA recommendation prototype "
        f"(IM/2021/049) — {date.today().strftime('%B %Y')}"
    )
    font(r, italic=True, size=11)

    note(
        doc,
        "This chapter replaces the earlier methodology wording that assumed BERT fine-tuning, "
        "large-scale manual gold annotation, and a transformer training split. It describes the "
        "system as implemented: Bright Data Amazon smartphone review corpus → preprocessing → "
        "lexicon-based ABSA (optional LLM API) → feature scores → weighted ranking + explanations "
        "→ offline technical evaluation against an Amazon star-rating baseline.",
    )

    # 3.1
    h(doc, "3.1 Chapter Introduction", 2)
    p(
        doc,
        "This chapter explains the methodology used to construct and evaluate the proposed "
        "smartphone recommendation artefact. The study uses a frozen offline corpus of publicly "
        "visible English-language Amazon smartphone reviews obtained through Bright Data’s "
        "Amazon cell-phone / smartphone review collection pathway and stored locally as a "
        "structured CSV file for reproducible analysis. Six review-derived features—camera, "
        "battery, display, performance, design, and price/value—are identified from the review "
        "text itself rather than copied from manufacturer specifications. Extracted feature "
        "opinions are classified as positive, negative, or neutral, aggregated into product-level "
        "feature scores, combined with user-defined feature weights, and converted into ranked "
        "recommendations with traceable explanations.",
    )
    p(
        doc,
        "The methodology covers dataset construction from the Bright Data Amazon scrape, "
        "preprocessing, sentence and clause segmentation, domain-lexicon aspect identification "
        "and sentiment classification (with an optional OpenAI-compatible large language model "
        "API as a secondary engine), feature-score calculation with shrinkage, weighted ranking, "
        "rule-based explanation generation, and offline technical evaluation. The primary ABSA "
        "engine used for the reported thesis runs is the offline domain lexicon. Local "
        "fine-tuning of BERT or similar transformer classifiers was not performed. No "
        "questionnaire or interview study with recruited human participants is required for the "
        "core evaluation path described in this chapter. The design is informed by the selected "
        "systematic literature evidence on smartphone review analysis, ABSA, sentiment-aware "
        "recommendation, and explainability (Soni et al., 2024; Aziz et al., 2024; Musto et al., "
        "2017; Haghighi & Moradian Zadeh, 2025).",
    )

    # 3.2
    h(doc, "3.2 Research Philosophy", 2)
    p(
        doc,
        "The study adopts pragmatism because methods are selected according to their usefulness "
        "in solving the smartphone decision-support problem. Amazon review text is treated as "
        "practical evidence, processed through natural language processing and aspect-level "
        "sentiment classification, transformed into numerical feature scores, and evaluated "
        "through measurable technical criteria.",
    )
    p(
        doc,
        "Pragmatism also supports iterative artefact development followed by controlled "
        "evaluation on a frozen review snapshot. Development work was used to refine aspect "
        "lexicon coverage, preprocessing rules, scoring parameters, and explanation templates. "
        "Once the final configuration was selected, evaluation was conducted on the frozen "
        "corpus and on the implemented ranking and explanation modules, rather than on a "
        "continually changing live Amazon scrape.",
    )

    # 3.3
    h(doc, "3.3 Research Approach and Type of Study", 2)
    p(
        doc,
        "The research is predominantly quantitative and computational. Review text is converted "
        "into aspect categories, sentiment labels, feature scores, user weights, recommendation "
        "scores, ranks, and performance indicators. The main findings therefore depend on "
        "measurable system outputs rather than manual thematic analysis of interview transcripts.",
    )
    p(
        doc,
        "It is an applied study because it produces a working prototype for smartphone selection. "
        "No human participants are recruited for questionnaires or interviews. Evaluation uses "
        "the prepared review corpus, predefined feature-weight profiles, an Amazon average-rating "
        "baseline, automated star–text consistency checks, and explanation faithfulness checks. "
        "Formal causal hypotheses are unnecessary because the purpose is artefact construction "
        "and performance evaluation rather than modelling latent social-science relationships.",
    )

    # 3.4
    h(doc, "3.4 Research Strategy", 2)
    p(
        doc,
        "An applied design-and-development strategy is combined with quantitative experimental "
        "evaluation. The artefact integrates data ingest, ABSA, feature scoring, preference "
        "weighting, recommendation, and explanation rather than treating sentiment classification "
        "as the final output.",
    )
    p(
        doc,
        "Development was iterative, but the evaluation corpus was frozen. Preprocessing rules, "
        "the aspect lexicon, aggregation parameters, and explanation templates were refined "
        "during implementation. Final descriptive statistics, ranking comparisons, and automated "
        "checks were then produced from the frozen SQLite research database and the FastAPI "
        "research interface.",
    )

    # 3.5
    h(doc, "3.5 Research Design and Time Horizon", 2)
    p(
        doc,
        "The design is sequential and modular: each stage produces a defined output for the next "
        "stage, from Bright Data Amazon review collection to offline system evaluation. "
        "Figure 3.1 presents this workflow.",
    )
    note(
        doc,
        "Figure 3.1: Research methodology workflow (as built). "
        "Bright Data Amazon smartphone reviews (CSV) → Step 1 Preprocess → Step 2 Segment → "
        "Steps 3–4 Lexicon ABSA (optional LLM) → Step 5 Aspect–sentiment store → Step 6 Feature "
        "scores → Weighted ranking + explanations → Offline evaluation vs Amazon star baseline. "
        "Source: Developed by the researcher.",
    )
    p(
        doc,
        "The time horizon is cross-sectional. Reviews were collected during a documented period "
        "through Bright Data and frozen as a versioned offline dataset (local CSV and SQLite "
        "database), ensuring that experiments use the same evidence and can be reproduced.",
    )

    table(
        doc,
        ["Methodological component", "Adopted choice", "Reason for the choice"],
        [
            [
                "Research philosophy",
                "Pragmatism",
                "Focuses on practical development and evaluation of a useful decision-support artefact.",
            ],
            [
                "Research approach",
                "Predominantly quantitative and computational",
                "Transforms review text into measurable labels, scores, rankings, and evaluation indicators.",
            ],
            [
                "Research type",
                "Applied research",
                "Addresses a practical smartphone-selection problem through a working prototype.",
            ],
            [
                "Research strategy",
                "Design and development with experimental evaluation",
                "Builds an integrated system and evaluates components using controlled criteria.",
            ],
            [
                "Time horizon",
                "Cross-sectional",
                "Uses a fixed snapshot of Amazon reviews obtained via Bright Data.",
            ],
            [
                "Main data source",
                "Secondary user-generated Amazon reviews (Bright Data collection)",
                "Reviews were originally created for product feedback, not specifically for this research.",
            ],
            [
                "Human participants",
                "Not recruited",
                "Core evaluation uses the review corpus, baselines, weight profiles, and automated checks.",
            ],
            [
                "Primary ABSA technique",
                "Domain lexicon (optional LLM API)",
                "Supports offline reproducibility without local BERT fine-tuning.",
            ],
        ],
        caption="Table 3.1: Overall methodological profile (as built)",
    )

    table(
        doc,
        ["Research objective", "Methodological procedure", "Expected output"],
        [
            [
                "RO1: Construct a reliable Amazon smartphone review dataset",
                "Bright Data Amazon cell-phone/smartphone review collection, local CSV storage, eligibility screening, English filtering, deduplication, and versioned SQLite storage",
                "Clean, traceable smartphone review dataset",
            ],
            [
                "RO2: Identify smartphone features and classify aspect sentiment",
                "Closed six-aspect taxonomy, domain lexicon matching with negation/contrast handling; optional LLM JSON annotation",
                "Aspect–sentiment records for camera, battery, display, performance, design, and price/value",
            ],
            [
                "RO3: Generate product-level feature performance scores",
                "Aggregate positive/neutral/negative mentions; compute shrunk feature scores on [0, 1]",
                "Comparable feature scores for each smartphone–aspect pair",
            ],
            [
                "RO4: Rank smartphones according to user priorities",
                "Normalise feature weights and compute a weighted recommendation score; compare with Amazon star baseline",
                "Personalised ranked list of smartphones",
            ],
            [
                "RO5: Generate transparent recommendation explanations",
                "Rule-based templates using feature scores, weights, contributions, and weaknesses",
                "Traceable natural-language explanation for each recommendation",
            ],
        ],
        caption="Table 3.2: Alignment between research objectives and methodological procedures (as built)",
    )

    # 3.6
    h(doc, "3.6 Population, Units of Analysis, and Sample", 2)
    h(doc, "3.6.1 Data Population", 3)
    p(
        doc,
        "The target data population consists of publicly accessible English-language reviews "
        "associated with smartphone listings on Amazon, collected through Bright Data’s Amazon "
        "cell-phone / smartphone review data pathway during the documented collection period. "
        "Because Amazon review content changes over time, the study uses a fixed and documented "
        "snapshot rather than attempting to include every review ever published.",
    )
    p(
        doc,
        "The reviews are treated as secondary user-generated data because they were originally "
        "posted by customers for product feedback and were not created specifically for this "
        "research. Reviewer names, profile details, and other personally identifying information "
        "are not required for the analysis and are not retained in the research database used for "
        "scoring and recommendation.",
    )

    h(doc, "3.6.2 Units of Analysis", 3)
    p(
        doc,
        "The primary unit of analysis is the aspect-level sentence or clause extracted from each "
        "smartphone review. Reviews are divided into sentences and, where contrast markers appear, "
        "into separate clauses so that different feature opinions are analysed independently.",
    )
    p(
        doc,
        "For example, the statement “The camera is excellent, but the battery drains quickly” is "
        "represented as two analytical records: Camera – positive; Battery – negative. The "
        "smartphone product is the aggregation unit because all aspect-level sentiment results "
        "for a product are combined to calculate feature performance scores for camera, battery, "
        "display, performance, design, and price.",
    )

    h(doc, "3.6.3 Sample (Actual Corpus)", 3)
    p(
        doc,
        "The computational sample is the Bright Data–sourced Amazon smartphone review file "
        "stored locally as full_reviews.csv and imported into SQLite. Unlike the earlier planning "
        "estimate of approximately 15 models with 200 reviews each, the implemented corpus is "
        "broader in product coverage and thinner per phone, which is reported honestly below.",
    )

    table(
        doc,
        ["Stage", "Unit", "Count"],
        [
            ["Raw Bright Data / local CSV snapshot", "Review rows / unique ASINs", "4,092 / 227"],
            ["After dropping manuals / guide-book listings", "Review rows / ASINs", "4,080 / 224"],
            ["Loaded into research database", "Phones / reviews", "224 / 4,080"],
            ["After Step 1 preprocessing (kept English usable reviews)", "Reviews", "3,389"],
            ["Excluded in Step 1", "Reviews", "691"],
            ["Step 2 sentence segmentation", "Sentences", "31,685"],
            ["Steps 3–4 ABSA output", "Aspect–sentiment records", "13,914"],
            ["Step 6 feature-score table", "Phones with at least one aspect score", "221"],
        ],
        caption="Table 3.3: Actual dataset funnel used in the study",
    )
    p(
        doc,
        "Of the 691 excluded reviews, 663 were removed by English language filtering; the "
        "remainder were removed for all-caps, no alphabetic content, extreme shortness, or "
        "generic-only text. Kept reviews per phone after preprocessing ranged from 1 to 24 "
        "(mean approximately 15.2). This distribution differs from the earlier planning target "
        "of ≥200 reviews per model; all accessible eligible reviews in the frozen Bright Data "
        "snapshot were retained rather than discarding products with fewer reviews.",
    )
    p(
        doc,
        "No large-scale manual gold annotation of 2,400 clauses and no 70:15:15 BERT training "
        "split were required for the as-built primary pipeline. Optional gold-label evaluation "
        "remains available as a future extension, not as a prerequisite for the reported system.",
    )

    # 3.7
    h(doc, "3.7 Sampling Strategy and Eligibility Criteria", 2)
    p(
        doc,
        "Smartphone listings present in the Bright Data Amazon cell-phone dataset were retained "
        "when they corresponded to genuine smartphone products. Non-phone listings such as user "
        "manuals and guide books were removed by title-pattern screening before analysis. "
        "Duplicate configurations were avoided unless they represented distinct ASINs with "
        "distinct review pools.",
    )
    p(
        doc,
        "Reviews were screened by predefined criteria. Public English text linked to an eligible "
        "phone was retained, while empty, extremely short, non-English, inaccessible, and exact "
        "duplicate records were removed during ingest and preprocessing.",
    )

    table(
        doc,
        ["Code", "Inclusion criteria", "Code", "Exclusion criteria"],
        [
            [
                "IC1",
                "Publicly visible Amazon smartphone reviews in the Bright Data snapshot",
                "EC1",
                "Accessories, manuals/guide books, tablets, wearables, or unrelated products",
            ],
            [
                "IC2",
                "Reviews written in English",
                "EC2",
                "Non-English reviews in the main implementation",
            ],
            [
                "IC3",
                "Reviews containing meaningful textual content",
                "EC3",
                "Empty, rating-only, or extremely short comments without usable opinion evidence",
            ],
            [
                "IC4",
                "Reviews collected within the documented Bright Data / local CSV procedure",
                "EC4",
                "Login-protected or inaccessible content not present in the snapshot",
            ],
            [
                "IC5",
                "Reviews that can contribute aspect evidence after ABSA",
                "EC5",
                "Records with no detectable relationship to the six selected aspects after analysis",
            ],
            [
                "IC6",
                "Unique review records after exact duplicate screening",
                "EC6",
                "Exact duplicates likely to inflate sentiment counts",
            ],
            [
                "IC7",
                "Reviews whose product identity can be verified through an ASIN",
                "EC7",
                "Records that cannot be reliably linked to a smartphone model",
            ],
        ],
        caption="Table 3.4: Inclusion and exclusion criteria for the Amazon review dataset",
    )

    # 3.8
    h(doc, "3.8 Data Collection Procedure", 2)
    h(doc, "3.8.1 Source and Collection Mode", 3)
    p(
        doc,
        "The primary data source is Amazon smartphone / cell-phone review content collected "
        "via Bright Data and delivered as a structured dataset. The local working file used by "
        "the prototype is data/full_reviews.csv. This path was selected to obtain a stable, "
        "reproducible offline corpus without performing live scraping at recommendation time.",
    )
    p(
        doc,
        "An optional Playwright-based live Amazon scraper remains in the codebase for "
        "supplementary collection experiments, but it is not required for the main thesis "
        "workflow. Camera, battery, display, performance, design, and price scores are derived "
        "later from opinions expressed in the reviews, not copied from product specifications.",
    )

    h(doc, "3.8.2 Collection and Ingest Steps", 3)
    for item in [
        "Obtain the Bright Data Amazon cell-phone / smartphone review export for the documented collection window.",
        "Store the unmodified export and create the local analysis file full_reviews.csv.",
        "Validate required fields (ASIN, product title, review text) and drop non-phone manual/guide listings.",
        "Import phones, prices, ratings metadata, and reviews into SQLite using the research ingest command.",
        "Run Step 1 preprocessing (English filter, spam/short-text controls, normalisation, de-duplication).",
        "Freeze the database snapshot used for ABSA, scoring, ranking, and evaluation.",
    ]:
        bullet(doc, item)

    table(
        doc,
        ["Data field", "Description", "Use in the study"],
        [
            ["Product identifier / ASIN", "Stable Amazon listing identifier", "Product grouping and traceability"],
            ["Product name and brand", "Displayed model and manufacturer", "Recommendation display and description"],
            ["Product URL", "Source listing URL where available", "Audit trail"],
            ["Listing price at collection", "Displayed price where available", "Optional budget filtering; not a sentiment label"],
            ["Average product rating", "Displayed aggregate Amazon rating", "Star-rating recommendation baseline"],
            ["Review text", "Main written review content", "Primary input to preprocessing and ABSA"],
            ["Review star rating", "Rating attached to the review", "Consistency checks / star–text agreement"],
            ["Image URL / metadata", "Listing media where available", "UI display only"],
        ],
        caption="Table 3.5: Data fields used from the Bright Data Amazon snapshot",
    )

    h(doc, "3.8.3 Data Storage and Security", 3)
    p(
        doc,
        "Raw CSV records are preserved separately from the cleaned SQLite research database. "
        "Aspect records, feature scores, and recommendation feedback (if any pilot logs are "
        "created during demonstration) are stored in versioned local files under the project "
        "data directory. Reviewer names and profile links are not required and are not retained "
        "for analysis outputs.",
    )

    # 3.9
    h(doc, "3.9 Data Preparation", 2)
    h(doc, "3.9.1 Preprocessing (Step 1)", 3)
    p(
        doc,
        "Preprocessing removes technical noise while preserving language needed for sentiment. "
        "Negations, intensifiers, and meaningful opinion terms are retained. Core steps include "
        "record validation, exact duplicate screening, Unicode and whitespace normalisation, "
        "HTML remnant removal, spam heuristics, and English language filtering using langdetect.",
    )

    table(
        doc,
        ["Procedure", "Implementation", "Output"],
        [
            ["Record validation", "Confirm ASIN, review text, and product linkage", "Traceable raw record"],
            ["Non-phone screening", "Drop manuals/guide-book title patterns", "Smartphone-only rows"],
            ["Exact duplicate control", "Content hashing / de-duplication during ingest", "Reduced repeated inflation"],
            ["Text normalisation", "HTML strip, Unicode NFKC, whitespace collapse", "Clean contextual text"],
            ["Language filtering", "Retain English reviews (langdetect)", "Consistent language scope"],
            ["Spam / quality filters", "Too short, all-caps, no-alpha, generic-only checks", "Usable opinion text"],
            ["Sentence segmentation", "pysbd sentence boundaries (Step 2)", "Sentence-level records"],
            ["Clause handling", "Contrast markers such as but / however in lexicon ABSA", "Reduced cross-clause sentiment leakage"],
        ],
        caption="Table 3.6: Data preparation procedures (as built)",
    )

    h(doc, "3.9.2 Aspect Taxonomy (No Large-Scale Gold Labelling Required)", 3)
    p(
        doc,
        "The system uses a closed annotation scheme for six aspects and three sentiments. In the "
        "as-built primary path, labels are produced automatically by the domain lexicon engine "
        "(or optionally by an LLM API). Large-scale researcher gold labelling for BERT training "
        "was not conducted. The scheme below defines the operational meaning of each label.",
    )

    table(
        doc,
        ["Element", "Allowed value", "Decision rule"],
        [
            [
                "Aspect category",
                "Camera; Battery; Display; Performance; Design; Price/Value",
                "Assign every detected opinion to the smartphone feature it evaluates.",
            ],
            [
                "Positive sentiment",
                "Positive",
                "The clause expresses satisfaction, strength, benefit, or favourable value.",
            ],
            [
                "Negative sentiment",
                "Negative",
                "The clause expresses dissatisfaction, weakness, failure, or unfavourable value.",
            ],
            [
                "Neutral sentiment",
                "Neutral",
                "The aspect is mentioned factually without a clear favourable or unfavourable judgement.",
            ],
            [
                "Multiple aspects",
                "Multiple records",
                "Create one record per aspect; do not force the whole sentence into a single label.",
            ],
            [
                "Opposing opinions",
                "Clause-level handling",
                "Separate opinions at contrast markers; do not use “mixed” as a fourth class.",
            ],
        ],
        caption="Table 3.7: Aspect and sentiment scheme",
    )

    # 3.10
    h(doc, "3.10 System Design and Development", 2)
    h(doc, "3.10.1 Aspect Identification", 3)
    p(
        doc,
        "Aspect identification uses a versioned domain lexicon of explicit terms and synonyms "
        "for camera, battery, display, performance, design, and price/value, implemented in "
        "app/nlp/aspects.py and applied by the lexicon ABSA engine in app/nlp/absa.py. Longer "
        "phrases are matched before shorter forms to reduce false splits. The closed aspect "
        "scheme supports transparency, although highly implicit aspects without recognisable "
        "cues may remain undetected.",
    )

    table(
        doc,
        ["Aspect", "Indicative terms", "Operational meaning"],
        [
            ["Camera", "camera, photo, selfie, video, lens, zoom, night mode", "Image/video quality and shooting experience"],
            ["Battery", "battery, charging, drain, backup, fast charging", "Endurance and power behaviour"],
            ["Display", "display, screen, brightness, AMOLED, refresh rate", "Visual quality and touch experience"],
            ["Performance", "speed, lag, gaming, heat, processor, RAM", "Speed, stability, thermals, multitasking"],
            ["Design", "design, build, look, weight, grip, premium feel", "Appearance, ergonomics, materials"],
            ["Price/Value", "price, cost, value, expensive, worth, budget", "Affordability and value for money in reviews"],
        ],
        caption="Table 3.8: Smartphone aspect categories and indicative terms",
    )

    h(doc, "3.10.2 Aspect-Level Sentiment Classification", 3)
    p(
        doc,
        "Primary technique: domain lexicon sentiment scoring with negation windows, "
        "intensifiers/diminishers, and contrast-marker handling. Each detected aspect mention "
        "is labelled positive, negative, or neutral and stored with method provenance "
        "(lexicon or lexicon_fallback).",
    )
    p(
        doc,
        "Optional secondary technique: an OpenAI-compatible LLM API can annotate batches of "
        "sentences as JSON when an API key is configured (ABSA_ENGINE=llm or auto). This is not "
        "equivalent to local BERT fine-tuning, does not use a labelled 70:15:15 training split, "
        "and was not required for the reproducible offline thesis runs. The reported main "
        "configuration is ABSA_ENGINE=lexicon.",
    )
    p(
        doc,
        "A classical TF-IDF–SVM or fine-tuned BERT classifier was not trained in the as-built "
        "system. Instead, technical validation emphasises corpus coverage, star–text agreement, "
        "ranking behaviour under fixed weight profiles, and explanation faithfulness.",
    )

    h(doc, "3.10.3 Feature Score Generation", 3)
    p(
        doc,
        "Aspect records are aggregated by smartphone and feature. Let P, N, and U represent "
        "positive, negative, and neutral mention counts for product p and feature a. The "
        "implemented raw score is:",
    )
    note(doc, "raw_score(p,a) = (P + 0.5 × U) / (P + N + U)   ∈ [0, 1]")
    p(
        doc,
        "Neutral mentions therefore contribute partially rather than being ignored or treated "
        "as full positives. The stored ranking score additionally applies Bayesian shrinkage "
        "toward the corpus mean for that aspect:",
    )
    note(
        doc,
        "score(p,a) = (mentions × raw_score + k × corpus_mean) / (mentions + k), with k = 5 by default.",
    )
    p(
        doc,
        "Positive, negative, and neutral counts are retained with each score so that evidence "
        "volume remains visible in the interface. Sparse aspects are not over-trusted. This "
        "formulation differs from the earlier proposal wording that used only the percentage of "
        "positive mentions × 100; the as-built formula is the one implemented in "
        "app/nlp/aggregate.py and used by the recommender.",
    )

    h(doc, "3.10.4 User Preference Weighting and Recommendation Ranking", 3)
    p(
        doc,
        "Users assign importance values (0–10 marks) to the six features in the research UI. "
        "The system normalises them so that they sum to one. If all entries are zero, equal "
        "weights are applied. For each eligible product:",
    )
    note(doc, "FinalScore(p) = Σ [NormalizedWeight(a) × score(p,a)]")
    p(
        doc,
        "Missing aspects are imputed with the candidate-set mean and flagged as imputed rather "
        "than forced to zero. Coverage reports the share of requested weight backed by real "
        "mentions. Optional min/max USD filters constrain list price. Smartphones are sorted "
        "from highest to lowest FinalScore. A baseline method ranks by Amazon average star "
        "rating only (site_rating). This explicit content-based design uses visible review "
        "evidence and stated preferences rather than hidden latent collaborative profiles "
        "(Musto et al., 2017; Haghighi & Moradian Zadeh, 2025).",
    )

    table(
        doc,
        ["Feature", "Example feature score (0–1)", "Normalised weight", "Weighted contribution"],
        [
            ["Camera", "0.84", "0.30", "0.252"],
            ["Battery", "0.78", "0.25", "0.195"],
            ["Display", "0.82", "0.15", "0.123"],
            ["Performance", "0.75", "0.15", "0.113"],
            ["Design", "0.70", "0.05", "0.035"],
            ["Price/Value", "0.80", "0.10", "0.080"],
            ["Final recommendation score", "—", "1.00", "0.798"],
        ],
        caption="Table 3.9: Example of weighted recommendation scoring (illustrative)",
    )

    h(doc, "3.10.5 Explanation Generation", 3)
    p(
        doc,
        "Explanations are rule-based to prevent unsupported statements. For each recommendation, "
        "the module reports the strongest weighted contributions from review-derived feature "
        "scores and notes an important weak feature when applicable. Baseline explanations "
        "explicitly state that Amazon star rating was used. Explanations are generated from the "
        "same calculation that produced the rank (app/services/recommender.py).",
    )

    # 3.11
    h(doc, "3.11 Operationalization of the Framework", 2)
    note(
        doc,
        "Figure 3.2: Operationalization map — Review text → Aspect + Sentiment → Feature score "
        "→ User weight → Weighted contribution → Final score / Rank → Explanation; Baseline = "
        "Amazon average rating.",
    )

    table(
        doc,
        ["Variable / output", "Role", "Operational definition", "Calculation / source"],
        [
            ["Review text", "Raw evidence", "Written customer opinion linked to a smartphone ASIN", "Bright Data CSV → SQLite"],
            ["Aspect category", "Intermediate", "Feature discussed in a sentence/clause", "Lexicon (optional LLM)"],
            ["Aspect sentiment", "Intermediate", "Polarity toward the aspect", "Positive / negative / neutral"],
            ["Feature mention count", "Evidence strength", "Valid aspect records for a product–feature pair", "P + N + U"],
            ["Feature score", "Input 1", "Shrunk review-based feature performance", "(P + 0.5U)/total with shrinkage"],
            ["User feature weight", "Input 2", "Relative importance assigned in the UI", "Normalised 0–1 weights"],
            ["Weighted contribution", "Intermediate", "Contribution of one feature to the product score", "score × weight"],
            ["Final recommendation score", "Output", "Fit of a phone to the user’s priorities", "Sum of contributions"],
            ["Recommendation rank", "Output", "Position after sorting", "Descending FinalScore"],
            ["Explanation output", "User-facing output", "Plain-language strengths/trade-offs", "Rule-based templates"],
            ["Baseline rating score", "Comparison", "Traditional aggregate rating", "Amazon site_rating"],
        ],
        caption="Table 3.10: Operationalization of the main variables and outputs (as built)",
    )

    # 3.12
    h(doc, "3.12 Data Analysis and Evaluation", 2)
    h(doc, "3.12.1 Descriptive Analysis", 3)
    p(
        doc,
        "Descriptive analysis reports products, reviews collected and retained, sentences, "
        "aspect–sentiment frequencies, and per-aspect coverage. Table 3.3 and the ABSA "
        "distributions in Chapter 5 establish the composition of the evidence before ranking "
        "results are interpreted.",
    )

    h(doc, "3.12.2 ABSA Technical Checks (As-Built Substitute for Gold Macro-F1)", 3)
    p(
        doc,
        "Because large-scale gold annotation and BERT test-set Macro-F1 were not part of the "
        "primary as-built path, ABSA quality is assessed through: (i) descriptive coverage of "
        "all six aspects; (ii) star–text agreement between review polarity and Amazon review "
        "stars as a sanity check; and (iii) qualitative inspection of extracted sentences in "
        "the Phones / detail views. Optional gold-label Macro-F1 remains a future extension.",
    )

    h(doc, "3.12.3 Recommendation Evaluation", 3)
    p(
        doc,
        "Recommendation evaluation is conducted offline on the frozen corpus. Predefined "
        "feature-weight profiles (balanced, camera-focused, battery-focused, performance-focused, "
        "design-focused, and price/value-focused) are applied to the feature-score database. "
        "The proposed weighted ranking is compared with the Amazon average-rating baseline "
        "under the same product pool and filters. Differences in Top-N membership and order are "
        "examined as evidence that feature priorities change recommendations relative to stars.",
    )
    p(
        doc,
        "Where a researcher demonstration uses the prototype’s rating panel, graded 1–5 ratings "
        "may additionally support NDCG@3 and Spearman summaries on the Evaluation tab. Those "
        "ratings are optional demonstration / pilot logs rather than a formal recruited user "
        "study and are not required for the core methodological claim of this chapter.",
    )

    h(doc, "3.12.4 Explanation Transparency and Faithfulness", 3)
    p(
        doc,
        "Explanation evaluation is automated and does not involve participants. For each "
        "recommended smartphone, the explanation must identify contributions consistent with "
        "the computed weights and scores and must not invent unsupported reasons. Numerical "
        "consistency checks compare displayed scores with stored aspect_scores values. Coverage "
        "is the percentage of recommendations for which an explanation is generated.",
    )

    table(
        doc,
        ["Component", "Metric / check used (as built)", "Purpose"],
        [
            ["Corpus preparation", "Funnel counts and exclusion reasons", "Documents retained evidence after preprocessing"],
            ["ABSA coverage", "Aspect and sentiment frequency tables", "Shows six-feature evidence volume"],
            ["ABSA sanity check", "Star–text agreement rate", "Checks polarity consistency with review stars"],
            ["Recommendation", "Proposed vs Amazon star baseline Top-N comparison", "Shows effect of user priorities"],
            ["Recommendation (optional pilot)", "NDCG@3 and Spearman from demo ratings", "Optional ranking-quality summary"],
            ["Explanation", "Faithfulness / numerical consistency checks", "Prevents unsupported explanations"],
            ["Robustness", "Multiple weight profiles and budget filters", "Checks behaviour under different preferences"],
        ],
        caption="Table 3.11: Evaluation metrics and checks used in the study (as built)",
    )

    # 3.13
    h(doc, "3.13 Reliability, Validity, and Quality Assurance", 2)
    bullet(doc, "Dataset reliability: source CSV preservation, SQLite freeze, duplicate and language controls.")
    bullet(doc, "Method reliability: closed aspect taxonomy, versioned lexicon, documented scoring equations.")
    bullet(doc, "Internal validity: offline ranking uses stored scores; baseline comparison uses the same candidate pool.")
    bullet(doc, "Construct validity: six review-derived features defined operationally; explanations bound to computed scores.")
    bullet(doc, "Robustness: multiple weight profiles, optional budget filters, and shrinkage for sparse mentions.")

    table(
        doc,
        ["Quality area", "Control measure", "Expected benefit"],
        [
            ["Source traceability", "Retain Bright Data/local CSV, ASIN, and ingest logs", "Supports auditability"],
            ["Raw-data preservation", "Keep immutable CSV separate from cleaned DB", "Prevents undocumented alteration"],
            ["Duplicate / language control", "De-duplication and English filter", "Reduces noise and inflation"],
            ["ABSA transparency", "Lexicon rules + stored aspect_sentiment rows", "Makes labels inspectable"],
            ["Scoring transparency", "Publish equations, counts, and shrinkage parameter", "Makes ranking logic auditable"],
            ["Explanation faithfulness", "Generate text only from computed evidence", "Prevents hallucinated reasons"],
            ["Baseline comparison", "Amazon star ranking under same filters", "Interpretable comparator"],
        ],
        caption="Table 3.12: Quality assurance measures",
    )

    # 3.14
    h(doc, "3.14 Ethical Considerations", 2)
    p(
        doc,
        "The study uses publicly oriented Amazon review content obtained through Bright Data "
        "and processed offline for academic analysis. No private account data or payment-protected "
        "content are required for the main workflow. Reviewer names and profile links are excluded "
        "from analysis outputs. Because no human participants are recruited for questionnaires or "
        "interviews, informed-consent procedures for surveys are not applicable. The prototype is "
        "a research decision-support tool, not a commercial endorsement of any product.",
    )

    # 3.15
    h(doc, "3.15 Tools and Technologies", 2)
    table(
        doc,
        ["Component", "Tool or technology", "Purpose"],
        [
            ["Programming language", "Python 3.12+", "Pipeline, scoring, API, evaluation"],
            ["Data source", "Bright Data Amazon cell-phone / smartphone reviews → local CSV", "Frozen offline corpus"],
            ["Data handling", "Pandas, CSV, SQLite, SQLAlchemy 2.0", "Ingest, storage, aggregation"],
            ["Preprocessing / NLP", "langdetect, pysbd, domain lexicon", "Cleaning, segmentation, ABSA"],
            ["Optional LLM client", "httpx + OpenAI-compatible API", "Secondary ABSA engine only"],
            ["API / server", "FastAPI, Uvicorn, Pydantic", "REST API and research UI hosting"],
            ["Recommendation engine", "Python scoring functions (recommender.py)", "Weights, ranking, explanations"],
            ["Prototype interface", "Custom HTML/CSS/JS UI at /ui/; Streamlit optional", "Priorities, ranks, evaluation views"],
            ["CLI", "Typer + Rich (run.py)", "prepare-csv, analyze, serve, stats, export"],
            ["Optional live scrape", "Playwright, BeautifulSoup4", "Supplementary collection only"],
            ["Version control", "Git + dated data files", "Reproducible experiment states"],
        ],
        caption="Table 3.13: Tools and technologies used in the study (as built)",
    )

    # 3.16
    h(doc, "3.16 Chapter Summary", 2)
    p(
        doc,
        "This chapter presented a pragmatic, predominantly quantitative design-and-development "
        "methodology for constructing and evaluating the proposed recommendation system as "
        "built. The process uses a frozen Bright Data Amazon smartphone review dataset, "
        "criterion-based product and review screening, controlled preprocessing, domain-lexicon "
        "ABSA with an optional LLM API, feature-score aggregation with shrinkage, normalised "
        "user weights, rule-based explanations, and offline technical evaluation against an "
        "Amazon star-rating baseline.",
    )
    p(
        doc,
        "Relative to the earlier submitted draft, the main methodological corrections are: "
        "(1) Bright Data offline CSV as the primary corpus source; (2) lexicon ABSA rather than "
        "fine-tuned BERT; (3) actual corpus sizes (3,389 kept reviews; 224 phones; 13,914 "
        "aspect–sentiment records); (4) FastAPI research stack with a custom UI; and "
        "(5) evaluation centred on offline technical checks and baseline comparison rather than "
        "gold-label Macro-F1 or a recruited participant study. These choices preserve the "
        "original research aim while documenting the implemented artefact accurately.",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
